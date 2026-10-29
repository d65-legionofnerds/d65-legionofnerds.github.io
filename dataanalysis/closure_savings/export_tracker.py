"""
Exports Scot's D65_Closure_Savings_Tracker.xlsx (Legion of Data Nerds) to the CSVs
the chart builder reads. Run once per tracker update:

    python export_tracker.py path/to/D65_Closure_Savings_Tracker.xlsx

Writes (data/):
  source_lines.csv   one row per object line per month (FY26 + FY27 reports)
  projections.csv    district's projected closure savings per item, by school
"""
import sys
from pathlib import Path
import openpyxl
import pandas as pd

OUT = Path(__file__).parent / "data"
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)

rows = [r for r in wb["Source Lines"].iter_rows(min_row=5, values_only=True) if r[0]]
src = pd.DataFrame(rows, columns=["category", "fund", "obj", "line_name", "fy", "fiscal_month",
                                  "month", "month_expended", "budget_in_report", "source_report", "note"])
src["fund"] = src["fund"].astype(str)
src["obj"] = src["obj"].astype(str)
src["fy"] = src["fy"].astype(int)
src["fiscal_month"] = src["fiscal_month"].astype(int)
src.to_csv(OUT / "source_lines.csv", index=False)

tr = wb["Tracker"]
proj = []
for r in range(6, 23):
    proj.append(dict(category=tr.cell(r, 1).value,
                     kingsley=tr.cell(r, 2).value, rhodes=tr.cell(r, 3).value,
                     combined=tr.cell(r, 4).value, proxy=tr.cell(r, 5).value,
                     kind=tr.cell(r, 6).value, signal=tr.cell(r, 16).value,
                     tracker_note=tr.cell(r, 19).value))
pd.DataFrame(proj).to_csv(OUT / "projections.csv", index=False)
print(len(src), "source lines;", len(proj), "projected items")
